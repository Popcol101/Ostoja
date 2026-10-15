"""Arkusz do poradnika: scenariusze testów, podsumowanie, rejestr błędów, notatki ze spotkań."""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

DIR = os.path.dirname(os.path.abspath(__file__))
scen = json.load(open(os.path.join(DIR, "scenariusze.json")))

BRAND, INK, MUTED, TINT, LINE = "3F5650", "1E2422", "5C6461", "E6EAE8", "D8DAD7"
FILL_IN = PatternFill("solid", fgColor="FFF2CC")  # komórki do wypełnienia
HEAD = PatternFill("solid", fgColor=TINT)
OKF = PatternFill("solid", fgColor="DCE8E2")
ERRF = PatternFill("solid", fgColor="F3E0DC")
thin = Side(style="thin", color=LINE)
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
F = lambda **k: Font(name="Arial", size=10, color=k.pop("color", INK), **k)
WRAP = Alignment(wrap_text=True, vertical="top")

wb = Workbook()

def header(ws, row, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), start=1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font = F(bold=True, color=BRAND)
        cell.fill = HEAD
        cell.border = BOX
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30

def body_cell(ws, r, c, v=None, fill_in=False):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = F()
    cell.alignment = WRAP
    cell.border = BOX
    if fill_in:
        cell.fill = FILL_IN
    return cell

# 1. Jak używać ---------------------------------------------------------------
ws = wb.active
ws.title = "Jak używać"
ws.column_dimensions["A"].width = 24
ws.column_dimensions["B"].width = 90
ws["A1"] = "Ostoja – testy wersji testowej"
ws["A1"].font = Font(name="Georgia", size=18, color=INK)
ws["A2"] = "Arkusz do poradnika „Ostoja-poradnik-wersji-testowej.docx”. Stan na 16 października 2026. Tylko dane fikcyjne."
ws["A2"].font = F(color=MUTED)
rows = [
    ("Zakładka", "Do czego służy"),
    ("Scenariusze", f"{len(scen)} scenariuszy z poradnika (rozdział 5). Przy każdym wybierz Wynik z listy (OK, Błąd, Nie sprawdzono), wpisz datę, tester i numer błędu. Wypełniasz tylko żółte kolumny."),
    ("Podsumowanie", "Liczy się samo: ile scenariuszy sprawdzono i ile ma błąd, w podziale na obszary. Nic tu nie wpisujesz."),
    ("Błędy", "Rejestr błędów i uwag. Jeden wiersz = jeden problem. Pierwszy wiersz to przykład – nadpisz go albo usuń."),
    ("Spotkania", "Notatki ze spotkań z zakładami: pięć stałych pytań z poradnika (rozdział 8). Pierwszy wiersz to przykład."),
    ("Kolor żółty", "Komórki do wypełnienia."),
    ("Waga błędu", "blokuje – nie da się dokończyć zadania; przeszkadza – da się, ale z trudem; drobiazg – wygląd, literówka."),
    ("Przed spotkaniem", "Przywróć dane testowe (GitHub, Actions, Wersja testowa – baza danych) i wyczyść kolumny Wynik, Data i Tester z poprzedniej rundy."),
]
for i, (a, b) in enumerate(rows, start=4):
    ca, cb = ws.cell(row=i, column=1, value=a), ws.cell(row=i, column=2, value=b)
    for c in (ca, cb):
        c.font = F(bold=(i == 4), color=BRAND if i == 4 else INK)
        c.alignment = WRAP
        c.border = BOX
        if i == 4:
            c.fill = HEAD
ws["A9"].fill = FILL_IN

# 2. Scenariusze --------------------------------------------------------------
sc = wb.create_sheet("Scenariusze")
cols = ["Nr", "Obszar", "Scenariusz", "Konto", "Co zrobić", "Oczekiwany wynik", "Urządzenie", "Wynik", "Data", "Tester", "Nr błędu / uwagi"]
header(sc, 1, cols, [6, 13, 26, 18, 48, 44, 12, 15, 12, 14, 26])
for r, s in enumerate(scen, start=2):
    vals = [s["id"], s["obszar"], s["nazwa"], s["konto"], s["kroki"], s["oczekiwane"], s["urzadzenie"]]
    for c, v in enumerate(vals, start=1):
        body_cell(sc, r, c, v)
    for c in range(8, 12):
        body_cell(sc, r, c, None, fill_in=True)
    sc.cell(row=r, column=3).font = F(bold=True)
    sc.cell(row=r, column=9).number_format = "DD.MM.YYYY"
last = len(scen) + 1
dv = DataValidation(type="list", formula1='"OK,Błąd,Nie sprawdzono"', allow_blank=True)
dv.error, dv.errorTitle = "Wybierz OK, Błąd albo Nie sprawdzono.", "Wynik"
sc.add_data_validation(dv)
dv.add(f"H2:H{last}")
dvd = DataValidation(type="date", operator="greaterThan", formula1="DATE(2026,1,1)", allow_blank=True)
dvd.error = "Wpisz datę, np. 16.10.2026."
sc.add_data_validation(dvd)
dvd.add(f"I2:I{last}")
sc.conditional_formatting.add(f"H2:H{last}", FormulaRule(formula=['$H2="OK"'], fill=OKF))
sc.conditional_formatting.add(f"H2:H{last}", FormulaRule(formula=['$H2="Błąd"'], fill=ERRF, font=Font(name="Arial", size=10, bold=True, color="8E3B33")))
sc.freeze_panes = "C2"
sc.auto_filter.ref = f"A1:K{last}"

# 3. Podsumowanie -------------------------------------------------------------
su = wb.create_sheet("Podsumowanie")
header(su, 1, ["Obszar", "Scenariusze", "OK", "Błąd", "Nie sprawdzono / puste", "Sprawdzone (%)"], [18, 13, 10, 10, 22, 16])
obszary = []
for s in scen:
    if s["obszar"] not in obszary:
        obszary.append(s["obszar"])
R = f"Scenariusze!$B$2:$B${last}"
H = f"Scenariusze!$H$2:$H${last}"
for i, o in enumerate(obszary, start=2):
    body_cell(su, i, 1, o)
    body_cell(su, i, 2, f'=COUNTIF({R},A{i})')
    body_cell(su, i, 3, f'=COUNTIFS({R},A{i},{H},"OK")')
    body_cell(su, i, 4, f'=COUNTIFS({R},A{i},{H},"Błąd")')
    body_cell(su, i, 5, f"=B{i}-C{i}-D{i}")
    c = body_cell(su, i, 6, f"=IF(B{i}=0,0,(C{i}+D{i})/B{i})")
    c.number_format = "0%"
t = len(obszary) + 2
body_cell(su, t, 1, "Razem").font = F(bold=True)
for col in "BCDE":
    cell = body_cell(su, t, " BCDE".index(col) + 1, f"=SUM({col}2:{col}{t-1})")
    cell.font = F(bold=True)
c = body_cell(su, t, 6, f"=IF(B{t}=0,0,(C{t}+D{t})/B{t})")
c.number_format = "0%"
c.font = F(bold=True)
su.cell(row=t + 2, column=1, value="Liczone z zakładki Scenariusze, kolumna Wynik. „Nie sprawdzono” i puste komórki liczą się razem.").font = F(color=MUTED)
su.cell(row=t + 3, column=1, value="Otwarte błędy (status inny niż sprawdzony lub odrzucony):").font = F(color=MUTED)
c = su.cell(row=t + 3, column=5, value='=COUNTA(Błędy!$A$2:$A$500)-COUNTIF(Błędy!$L$2:$L$500,"sprawdzony")-COUNTIF(Błędy!$L$2:$L$500,"odrzucony")')
c.font = F(bold=True)

# 4. Błędy ---------------------------------------------------------------------
bl = wb.create_sheet("Błędy")
cols = ["Nr", "Data", "Zgłaszający", "Scenariusz", "Adres strony", "Konto", "Urządzenie i przeglądarka",
        "Co robiłem (kliknięcia)", "Czego oczekiwałem", "Co się stało", "Waga", "Status", "Zrzut ekranu (link)"]
header(bl, 1, cols, [8, 11, 14, 11, 26, 16, 20, 34, 28, 28, 13, 14, 22])
example = ["B-001", None, "Ola (przykład)", "Z4", "/panel/zlecenia/… (karta Haliny Nowak)", "Anna Cichońska", "Android 15, Chrome",
           "Dodaj termin → Ceremonia, karawan 1, 18.10, 11:00 → Zapisz", "Komunikat o zajętym karawanie",
           "Termin zapisał się mimo kolizji", "przeszkadza", "nowy", "link do zrzutu w Dysku"]
for c, v in enumerate(example, start=1):
    cell = body_cell(bl, 2, c, v, fill_in=True)
    cell.font = F(italic=True, color=MUTED)
bl.cell(row=2, column=2, value="=DATE(2026,10,16)").number_format = "DD.MM.YYYY"
for r in range(3, 201):
    for c in range(1, 14):
        body_cell(bl, r, c, None, fill_in=c in (2, 3, 4, 8, 9, 10, 11, 12))
    bl.cell(row=r, column=2).number_format = "DD.MM.YYYY"
dvw = DataValidation(type="list", formula1='"blokuje,przeszkadza,drobiazg"', allow_blank=True)
dvs = DataValidation(type="list", formula1='"nowy,w naprawie,naprawiony,sprawdzony,odrzucony"', allow_blank=True)
for d in (dvw, dvs):
    bl.add_data_validation(d)
dvw.add("K2:K200")
dvs.add("L2:L200")
bl.conditional_formatting.add("K2:K200", FormulaRule(formula=['$K2="blokuje"'], fill=ERRF, font=Font(name="Arial", size=10, bold=True, color="8E3B33")))
bl.freeze_panes = "B2"
bl.auto_filter.ref = "A1:M200"

# 5. Spotkania -----------------------------------------------------------------
sp = wb.create_sheet("Spotkania")
qs = ["1. Co w panelu zajęło najwięcej czasu albo było niejasne?",
      "2. Czy 15 minut na odpowiedź jest realne o każdej porze, także w nocy?",
      "3. Czego brakuje w karcie zlecenia, żeby nie wracać do zeszytu?",
      "4. Czy cennik i porównanie cen są uczciwe wobec zakładu?",
      "5. Czy zapłaciłby 250 zł za zlecenie z podpisaną umową? Jeśli nie, ile i za co?"]
header(sp, 1, ["Zakład", "Miasto", "Data", "Rozmówca i rola", *qs, "Inne uwagi", "Następny krok"], [24, 12, 11, 20, 32, 32, 32, 32, 32, 28, 26])
sp.row_dimensions[1].height = 54
ex = ["Dom Pogrzebowy Cisza (przykład)", "Łódź", None, "Anna, właścicielka",
      "Dodawanie terminu: za dużo pól na telefonie.", "W dzień tak; w nocy 30 minut.", "Miejsce na numer chłodni i uwagi dla ekipy.",
      "Tak, jeśli dodatki nie są liczone w porównaniu.", "Tak, 250 zł za podpisaną umowę jest w porządku.", "Chce widzieć opinie rodzin.", "Pilot od 1 listopada"]
for c, v in enumerate(ex, start=1):
    cell = body_cell(sp, 2, c, v, fill_in=True)
    cell.font = F(italic=True, color=MUTED)
sp.cell(row=2, column=3, value="=DATE(2026,10,20)").number_format = "DD.MM.YYYY"
for r in range(3, 31):
    for c in range(1, 12):
        body_cell(sp, r, c, None, fill_in=True)
    sp.cell(row=r, column=3).number_format = "DD.MM.YYYY"
sp.freeze_panes = "B2"

for w in wb.worksheets:
    w.sheet_view.zoomScale = 100
    w.page_setup.orientation = "landscape"
    w.page_setup.paperSize = w.PAPERSIZE_A4
    w.page_setup.fitToWidth = 1
    w.page_setup.fitToHeight = 0
    w.sheet_properties.pageSetUpPr.fitToPage = True
    w.print_title_rows = "1:1" if w.title != "Jak używać" else None
    w.sheet_properties.tabColor = BRAND if w.title in ("Scenariusze", "Jak używać") else LINE
out = os.path.join(DIR, "Ostoja-scenariusze-testow.xlsx")
wb.save(out)
print(out)
